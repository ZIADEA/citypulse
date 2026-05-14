"""
reports_widget.py — Rapports v2.0
=================================
QSplitter 30/70 : catalogue catégories | configuration + aperçu WebEngine.
Génération en QThread + LoadingOverlay, historique reports_history, timer planif. 60s.
"""

import os
import json
import shutil
import tempfile
from datetime import datetime, date, timedelta

import re

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QComboBox,
    QFileDialog, QMessageBox, QFrame, QSplitter, QListWidget, QListWidgetItem,
    QStackedWidget, QDateEdit, QSpinBox, QLineEdit, QGroupBox,
    QCheckBox, QTimeEdit, QScrollArea, QTextEdit, QCompleter, QAbstractItemView,
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal, QTimer, QUrl, QTime
from PyQt6.QtGui import QFont

from ..database.db_manager import get_connection, log_action
from ..services.report_service import (
    ReportService,
    REPORTLAB_OK,
    OPENPYXL_OK,
)
from .loading_overlay import LoadingOverlay
from .help_dialog import show_help
from .lucide_icons import apply_action_button
from .toast import show_toast

from .webengine_support import HAS_WEB, QWebEngineView, QWebEngineSettings, WEBENGINE_FALLBACK_SHORT

try:
    from PyQt6.QtPdf import QPdfDocument
    from PyQt6.QtPdfWidgets import QPdfView
    HAS_PDF_VIEW = True
except ImportError:
    HAS_PDF_VIEW = False

C = {
    "bg": "#0D1B2A",
    "panel": "#112240",
    "border": "#1E3A5F",
    "accent": "#00D4FF",
    "text": "#E8F4FD",
    "muted": "#8899AA",
}


class _ReportWorker(QThread):
    finished_ok = pyqtSignal(str)
    failed = pyqtSignal(str)

    def __init__(self, fn, *args, **kwargs):
        super().__init__()
        self._fn = fn
        self._args = args
        self._kwargs = kwargs

    def run(self):
        try:
            path = self._fn(*self._args, **self._kwargs)
            self.finished_ok.emit(path)
        except Exception as e:
            self.failed.emit(str(e))


class _PreviewPane(QWidget):
    """Zone d'aperçu unifiée : QPdfView natif pour PDF, QWebEngineView pour HTML/XLSX."""

    def __init__(self, border_color: str, height: int = 380):
        super().__init__()
        self.setFixedHeight(height)
        self.setObjectName("previewPane")
        self.setStyleSheet(
            f"QWidget#previewPane{{border:1px solid {border_color};"
            "border-radius:6px;background:#0D1B2A;}}"
        )
        lo = QVBoxLayout(self)
        lo.setContentsMargins(0, 0, 0, 0)
        lo.setSpacing(0)

        self._pages = QStackedWidget(self)
        lo.addWidget(self._pages)

        # Placeholder vide
        self._placeholder = QLabel("— Aucun aperçu —")
        self._placeholder.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._placeholder.setStyleSheet("color:#8899AA;font-size:11px;")
        self._pages.addWidget(self._placeholder)  # index 0 toujours

        # QPdfView (PDF natif Qt 6) — index dynamique
        self._pdf_doc: "QPdfDocument | None" = None
        self._pdf_view: "QPdfView | None" = None
        self._pdf_view_idx: int = -1
        if HAS_PDF_VIEW:
            self._pdf_doc = QPdfDocument(self)
            self._pdf_view = QPdfView(self)
            self._pdf_view.setDocument(self._pdf_doc)
            try:
                self._pdf_view.setPageMode(QPdfView.PageMode.MultiPage)
            except AttributeError:
                pass
            self._pdf_view.setStyleSheet("background:#f5f5f5;")
            self._pdf_view_idx = self._pages.addWidget(self._pdf_view)

        # QWebEngineView (HTML / XLSX) — index dynamique
        self._web: "QWebEngineView | None" = None
        self._web_idx: int = -1
        if HAS_WEB:
            self._web = QWebEngineView(self)
            _s = self._web.settings()
            _s.setAttribute(QWebEngineSettings.WebAttribute.LocalContentCanAccessFileUrls, True)
            _s.setAttribute(QWebEngineSettings.WebAttribute.LocalContentCanAccessRemoteUrls, True)
            _s.setAttribute(QWebEngineSettings.WebAttribute.PluginsEnabled, True)
            try:
                _s.setAttribute(QWebEngineSettings.WebAttribute.PdfViewerEnabled, True)
            except AttributeError:
                pass
            self._web.setUrl(QUrl("about:blank"))
            self._web_idx = self._pages.addWidget(self._web)

    # ── API publique ─────────────────────────────────────────────────────────

    def show_pdf(self, path: str) -> bool:
        """Affiche un PDF via QPdfView (natif), sinon QWebEngineView en fallback."""
        if HAS_PDF_VIEW and self._pdf_doc is not None and self._pdf_view_idx >= 0:
            self._pdf_doc.close()
            self._pdf_doc.load(path)
            if self._pdf_doc.pageCount() > 0:
                self._pages.setCurrentIndex(self._pdf_view_idx)
                return True
        if HAS_WEB and self._web is not None and self._web_idx >= 0:
            self._web.setUrl(QUrl.fromLocalFile(os.path.abspath(path)))
            self._pages.setCurrentIndex(self._web_idx)
            return True
        return False

    def setHtml(self, html: str, base_url: "QUrl | None" = None):
        """Affiche du HTML (aperçu CGU, XLSX, etc.)."""
        if HAS_WEB and self._web is not None and self._web_idx >= 0:
            if base_url:
                self._web.setHtml(html, base_url)
            else:
                self._web.setHtml(html)
            self._pages.setCurrentIndex(self._web_idx)

    def setUrl(self, url: "QUrl"):
        """Navigation directe (rarement utilisée)."""
        if HAS_WEB and self._web is not None and self._web_idx >= 0:
            self._web.setUrl(url)
            self._pages.setCurrentIndex(self._web_idx)

    def reset(self):
        self._pages.setCurrentIndex(0)


class ReportsWidget(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self._service = ReportService()
        self._worker: _ReportWorker | None = None
        self._last_pdf_path: str | None = None
        self._last_pdf_default_name: str = "rapport.pdf"
        self._sched_last_kpi_day: str | None = None
        self._setup_ui()

        self._sched_timer = QTimer(self)
        self._sched_timer.setInterval(60_000)
        self._sched_timer.timeout.connect(self._on_schedule_tick)
        self._sched_timer.start()

    def _setup_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        header = QHBoxLayout()
        title = QLabel("Rapports & exports")
        title.setObjectName("heading")
        header.addWidget(title)
        header.addStretch()
        hb = QPushButton()
        hb.setFixedSize(28, 28)
        hb.setToolTip("Aide — Rapports et exports")
        hb.setCursor(Qt.CursorShape.PointingHandCursor)
        apply_action_button(hb, "help-circle", "#7FA8C0", "#1A2E4A", "#1A3A5C", 16)
        hb.clicked.connect(lambda: show_help(self, "reports"))
        header.addWidget(hb)
        root.addLayout(header)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.setChildrenCollapsible(False)
        splitter.setStyleSheet(f"QSplitter::handle{{background:{C['border']};width:3px;}}")

        # ── Gauche : catalogue ───────────────────────────────────────────
        left = QFrame()
        left.setFixedWidth(260)
        left.setStyleSheet(f"QFrame{{background:{C['panel']};border-right:1px solid {C['border']};}}")
        ll = QVBoxLayout(left)
        ll.setContentsMargins(8, 8, 8, 8)
        cat_lbl = QLabel("Catégories")
        cat_lbl.setStyleSheet(f"color:{C['muted']};font-size:11px;font-weight:600;")
        ll.addWidget(cat_lbl)

        lang_lbl = QLabel("Langue du rapport")
        lang_lbl.setStyleSheet(f"color:{C['muted']};font-size:10px;margin-top:6px;")
        ll.addWidget(lang_lbl)
        self._report_lang_combo = QComboBox()
        self._report_lang_combo.addItem("Français (FR)", "fr")
        self._report_lang_combo.addItem("English (EN)", "en")
        self._report_lang_combo.addItem("Español (ES)", "es")
        self._report_lang_combo.addItem("Deutsch (DE)", "de")
        self._report_lang_combo.addItem("العربية (AR)", "ar")
        self._report_lang_combo.setStyleSheet(
            f"QComboBox{{background:{C['bg']};color:{C['text']};border:1px solid {C['border']};"
            "border-radius:4px;padding:4px;font-size:11px;}}"
            f"QComboBox::drop-down{{border:none;}} QComboBox QAbstractItemView{{background:{C['bg']};color:{C['text']};}}"
        )
        ll.addWidget(self._report_lang_combo)

        self._catalog = QListWidget()
        self._catalog.setStyleSheet(
            f"QListWidget{{background:{C['bg']};border:1px solid {C['border']};border-radius:6px;}}"
            f"QListWidget::item{{padding:10px 8px;border-radius:4px;}}"
            f"QListWidget::item:selected{{background:{C['accent']};color:{C['bg']};}}"
        )
        self._catalog.setSpacing(2)

        categories = [
            ("Opérationnels", "ops"),
            ("Analytiques", "analytics"),
            ("Clients", "clients"),
            ("Transporteurs", "carriers"),
            ("Conformité", "compliance"),
            (" Documents légaux", "legal"),
            ("Exports", "exports"),
        ]
        for label, key in categories:
            it = QListWidgetItem(f"  {label}")
            it.setData(Qt.ItemDataRole.UserRole, key)
            it.setFont(QFont("Segoe UI", 11))
            self._catalog.addItem(it)

        self._catalog.currentItemChanged.connect(self._on_category_changed)
        ll.addWidget(self._catalog, 1)

        # Planification auto
        sched = QGroupBox("Planification auto (vérification 60s)")
        sched.setStyleSheet(
            f"QGroupBox{{font-size:11px;color:{C['muted']};border:1px solid {C['border']};"
            "border-radius:6px;margin-top:8px;padding-top:8px;}}"
            f"QGroupBox::title{{subcontrol-origin:margin;left:8px;padding:0 4px;}}"
        )
        sl = QVBoxLayout(sched)
        self._sched_enable = QCheckBox("Export KPI quotidien (PDF)")
        self._sched_enable.setStyleSheet(f"color:{C['text']};")
        sl.addWidget(self._sched_enable)
        dir_row = QHBoxLayout()
        dir_row.addWidget(QLabel("Dossier :"))
        self._sched_dir = QLineEdit()
        self._sched_dir.setPlaceholderText(os.path.expanduser("~/Documents/CityPulseReports"))
        self._sched_dir.setStyleSheet(
            f"QLineEdit{{background:{C['bg']};color:{C['text']};border:1px solid {C['border']};border-radius:4px;padding:4px;}}"
        )
        bdir = QPushButton("…")
        bdir.setFixedWidth(28)
        bdir.setToolTip("Choisir le dossier d'export planifié")
        bdir.clicked.connect(self._pick_sched_dir)
        dir_row.addWidget(self._sched_dir, 1)
        dir_row.addWidget(bdir)
        sl.addLayout(dir_row)
        tm_row = QHBoxLayout()
        tm_row.addWidget(QLabel("Heure :"))
        self._sched_time = QTimeEdit()
        self._sched_time.setDisplayFormat("HH:mm")
        self._sched_time.setTime(QTime(8, 0))
        self._sched_time.setStyleSheet(
            f"QTimeEdit{{background:{C['bg']};color:{C['text']};border:1px solid {C['border']};border-radius:4px;}}"
        )
        tm_row.addWidget(self._sched_time)
        tm_row.addStretch()
        sl.addLayout(tm_row)
        ll.addWidget(sched)

        # Connecter les changements → sauvegarde immédiate
        self._sched_enable.toggled.connect(self._save_sched_settings)
        self._sched_dir.textChanged.connect(self._save_sched_settings)
        self._sched_time.timeChanged.connect(self._save_sched_settings)

        splitter.addWidget(left)

        # ── Droite : tout le contenu dans un seul QScrollArea ───────────
        right_inner = QWidget()
        right_inner.setObjectName("rptRightInner")
        right_inner.setStyleSheet("QWidget#rptRightInner{background:transparent;}")
        rl = QVBoxLayout(right_inner)
        rl.setContentsMargins(12, 8, 12, 16)
        rl.setSpacing(14)

        self._stack = QStackedWidget()
        self._stack.setStyleSheet("background:transparent;")
        self._build_stacks()
        rl.addWidget(self._stack)

        hist_lbl = QLabel("Historique récent")
        hist_lbl.setStyleSheet(f"color:{C['muted']};font-size:11px;")
        rl.addWidget(hist_lbl)
        self._history_list = QListWidget()
        self._history_list.setMinimumHeight(100)
        self._history_list.setMaximumHeight(160)
        self._history_list.setStyleSheet(
            f"QListWidget{{background:{C['bg']};border:1px solid {C['border']};font-size:10px;}}"
        )
        self._history_list.itemDoubleClicked.connect(self._open_history_file)
        rl.addWidget(self._history_list)
        rl.addStretch()

        right_scroll = QScrollArea()
        right_scroll.setWidgetResizable(True)
        right_scroll.setWidget(right_inner)
        right_scroll.setStyleSheet(
            "QScrollArea{border:none;background:transparent;}"
            "QScrollArea>QWidget>QWidget{background:transparent;}"
        )

        splitter.addWidget(right_scroll)
        splitter.setSizes([280, 920])
        splitter.setStretchFactor(1, 1)
        root.addWidget(splitter, 1)

        self._overlay = LoadingOverlay(self)

        self._catalog.setCurrentRow(0)

    def resizeEvent(self, e):
        super().resizeEvent(e)
        if self._overlay.isVisible():
            self._overlay.setGeometry(self.rect())

    def _build_stacks(self):
        # 0 ops
        self._stack.addWidget(self._make_ops_panel())
        # 1 analytics
        self._stack.addWidget(self._make_analytics_panel())
        # 2 clients
        self._stack.addWidget(self._make_clients_panel())
        # 3 carriers
        self._stack.addWidget(self._make_carriers_panel())
        # 4 compliance
        self._stack.addWidget(self._make_compliance_panel())
        # 5 legal
        self._stack.addWidget(self._make_legal_panel())
        # 6 exports
        self._stack.addWidget(self._make_exports_panel())

    def _panel_shell(self, title: str) -> tuple[QWidget, QVBoxLayout]:
        w = QWidget()
        w.setObjectName("rptPanelInner")
        w.setStyleSheet("QWidget#rptPanelInner{background:transparent;}")
        lo = QVBoxLayout(w)
        lo.setContentsMargins(0, 0, 0, 8)
        lo.setSpacing(16)
        t = QLabel(title)
        t.setStyleSheet(f"font-size:13px;font-weight:700;color:{C['accent']};")
        lo.addWidget(t)
        return w, lo

    @staticmethod
    def _frm_row(label: str, widget, label_w: int = 150) -> "QHBoxLayout":
        row = QHBoxLayout()
        row.setSpacing(8)
        lbl = QLabel(label)
        lbl.setFixedWidth(label_w)
        lbl.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        lbl.setStyleSheet(f"color:{C['muted']};font-size:11px;border:none;background:transparent;")
        row.addWidget(lbl)
        row.addWidget(widget)
        row.addStretch()
        return row

    @staticmethod
    def _box(border_color: str) -> "tuple[QFrame, QVBoxLayout]":
        f = QFrame()
        f.setStyleSheet(
            f"QFrame#rptBox{{border:1px solid {border_color};border-radius:8px;background:transparent;}}"
        )
        f.setObjectName("rptBox")
        lo = QVBoxLayout(f)
        lo.setContentsMargins(12, 10, 12, 10)
        lo.setSpacing(8)
        return f, lo

    def _make_preview_block(self) -> tuple:
        """Returns (preview_pane, dl_btn, preview_header_lbl)."""
        header = QLabel("Aperçu")
        header.setStyleSheet(f"color:{C['muted']};font-size:11px;font-weight:600;")
        preview = _PreviewPane(C["border"], height=380)
        dl_btn = QPushButton("⬇  Télécharger le rapport")
        dl_btn.setObjectName("primaryBtn")
        dl_btn.setVisible(False)
        dl_btn.clicked.connect(self._download_last_report)
        return preview, dl_btn, header

    def _field_qss(self, kind: str = "input") -> str:
        base = f"border:1px solid {C['border']};border-radius:4px;padding:4px;background:{C['bg']};color:{C['text']};"
        if kind == "combo":
            return f"QComboBox{{{base}}} QComboBox::drop-down{{border:none;}} QComboBox QAbstractItemView{{background:{C['bg']};color:{C['text']};}}"
        if kind == "spin":
            return f"QSpinBox{{{base}}}"
        if kind == "date":
            return f"QDateEdit{{{base}}}"
        return f"QLineEdit{{{base}}}"

    def _make_ops_panel(self):
        w, lo = self._panel_shell("Rapports opérationnels")
        box, bl = self._box(C["border"])

        _cqss = self._field_qss("combo")
        self._op_route_combo = QComboBox()
        self._op_route_combo.setEditable(True)
        self._op_route_combo.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self._op_route_combo.setStyleSheet(_cqss)
        self._op_route_combo.lineEdit().setPlaceholderText("Date ou immatriculation véhicule…")

        self._op_date = QDateEdit()
        self._op_date.setCalendarPopup(True)
        self._op_date.setDate(datetime.today().date())
        self._op_date.setMaximumWidth(150)
        self._op_date.setStyleSheet(self._field_qss("date"))

        bl.addLayout(self._frm_row("Tournée :", self._op_route_combo))
        bl.addLayout(self._frm_row("Date flotte :", self._op_date))
        lo.addWidget(box)

        btn_row = QHBoxLayout()
        btn_row.setSpacing(6)
        b1 = QPushButton("Carnet chauffeur (PDF)")
        b1.setObjectName("primaryBtn")
        b1.clicked.connect(self._gen_roadbook)
        b2 = QPushButton("Rapport flotte jour (PDF)")
        b2.setObjectName("secondaryBtn")
        b2.clicked.connect(self._gen_fleet_daily)
        btn_row.addWidget(b1)
        btn_row.addWidget(b2)
        lo.addLayout(btn_row)

        self._ops_preview, self._ops_dl_btn, _ops_prev_hdr = self._make_preview_block()
        lo.addWidget(_ops_prev_hdr)
        lo.addWidget(self._ops_preview)
        lo.addWidget(self._ops_dl_btn)
        lo.addStretch()
        return w

    def _make_analytics_panel(self):
        w, lo = self._panel_shell("Rapports analytiques")
        box, bl = self._box(C["border"])

        self._an_kpi_start = QDateEdit()
        self._an_kpi_end = QDateEdit()
        for d in (self._an_kpi_start, self._an_kpi_end):
            d.setCalendarPopup(True)
            d.setDate(datetime.today().date())
            d.setMaximumWidth(150)
            d.setStyleSheet(self._field_qss("date"))

        self._an_kpi_fmt = QComboBox()
        self._an_kpi_fmt.addItems(["pdf", "xlsx"])
        self._an_kpi_fmt.setMaximumWidth(110)
        self._an_kpi_fmt.setStyleSheet(self._field_qss("combo"))

        _cqss = self._field_qss("combo")
        self._an_algo_pick = QComboBox()
        self._an_algo_pick.setEditable(True)
        self._an_algo_pick.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self._an_algo_pick.setStyleSheet(_cqss)
        self._an_algo_pick.lineEdit().setPlaceholderText("Chercher par algo, km, date…")
        self._an_algo_pick.currentIndexChanged.connect(self._on_algo_pick)

        # Liste des résultats sélectionnés (remplace le champ texte ID brut)
        self._an_algo_list = QListWidget()
        self._an_algo_list.setMaximumHeight(100)
        self._an_algo_list.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self._an_algo_list.setStyleSheet(
            f"QListWidget{{background:{C['bg']};color:{C['text']};"
            f"border:1px solid {C['border']};border-radius:4px;font-size:11px;}}"
            f"QListWidget::item{{padding:3px 6px;}}"
            f"QListWidget::item:selected{{background:{C['accent']};color:{C['bg']};}}"
        )
        algo_rm_btn = QPushButton("✕  Retirer la sélection")
        algo_rm_btn.setObjectName("ghostBtn")
        algo_rm_btn.setFixedHeight(24)
        algo_rm_btn.setStyleSheet(
            f"QPushButton{{font-size:10px;color:{C['muted']};border:none;background:transparent;text-align:left;padding:0 4px;}}"
            f"QPushButton:hover{{color:{C['text']};}}"
        )
        algo_rm_btn.clicked.connect(self._algo_list_remove)

        self._an_drv_days = QSpinBox()
        self._an_drv_days.setRange(1, 365)
        self._an_drv_days.setValue(30)
        self._an_drv_days.setMaximumWidth(100)
        self._an_drv_days.setStyleSheet(self._field_qss("spin"))

        self._an_drv_fmt = QComboBox()
        self._an_drv_fmt.addItems(["pdf", "xlsx"])
        self._an_drv_fmt.setMaximumWidth(110)
        self._an_drv_fmt.setStyleSheet(self._field_qss("combo"))

        bl.addLayout(self._frm_row("KPI — début :", self._an_kpi_start))
        bl.addLayout(self._frm_row("KPI — fin :", self._an_kpi_end))
        bl.addLayout(self._frm_row("Format KPI :", self._an_kpi_fmt))
        bl.addLayout(self._frm_row("Ajouter résultat :", self._an_algo_pick))

        # Ligne "Résultats à comparer" avec la liste + bouton retirer
        _algo_hdr = QHBoxLayout()
        _algo_lbl = QLabel("Résultats sélectionnés :")
        _algo_lbl.setFixedWidth(150)
        _algo_lbl.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignTop)
        _algo_lbl.setStyleSheet(f"color:{C['muted']};font-size:11px;border:none;background:transparent;")
        _algo_right = QVBoxLayout()
        _algo_right.setSpacing(3)
        _algo_right.addWidget(self._an_algo_list)
        _algo_right.addWidget(algo_rm_btn)
        _algo_hdr.addWidget(_algo_lbl)
        _algo_hdr.addLayout(_algo_right)
        _algo_hdr.addStretch()
        bl.addLayout(_algo_hdr)

        bl.addLayout(self._frm_row("Perf. (jours) :", self._an_drv_days))
        bl.addLayout(self._frm_row("Format perf. :", self._an_drv_fmt))
        lo.addWidget(box)

        r1 = QHBoxLayout()
        r1.setSpacing(6)
        for txt, slot, obj in [
            ("Rapport KPI", self._gen_kpi, "primaryBtn"),
            ("Performance chauffeurs", self._gen_drv_perf, "secondaryBtn"),
        ]:
            b = QPushButton(txt)
            b.setObjectName(obj)
            b.clicked.connect(slot)
            r1.addWidget(b)
        lo.addLayout(r1)

        self._an_preview, self._an_dl_btn, _an_prev_hdr = self._make_preview_block()
        lo.addWidget(_an_prev_hdr)
        lo.addWidget(self._an_preview)
        lo.addWidget(self._an_dl_btn)
        lo.addStretch()
        return w

    def _make_clients_panel(self):
        w, lo = self._panel_shell("Rapports clients")
        box, bl = self._box(C["border"])

        self._cl_combo = QComboBox()
        self._cl_combo.setEditable(True)
        self._cl_combo.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self._cl_combo.setStyleSheet(self._field_qss("combo"))
        self._cl_combo.lineEdit().setPlaceholderText("Nom du client ou entreprise…")

        bl.addLayout(self._frm_row("Client :", self._cl_combo))
        lo.addWidget(box)

        b = QPushButton("Générer fiche client (PDF)")
        b.setObjectName("primaryBtn")
        b.clicked.connect(self._gen_client)
        lo.addWidget(b)

        self._cl_preview, self._cl_dl_btn, _cl_prev_hdr = self._make_preview_block()
        lo.addWidget(_cl_prev_hdr)
        lo.addWidget(self._cl_preview)
        lo.addWidget(self._cl_dl_btn)
        lo.addStretch()
        return w

    def _make_carriers_panel(self):
        w, lo = self._panel_shell("Transporteurs")
        box, bl = self._box(C["border"])

        self._car_combo = QComboBox()
        self._car_combo.setEditable(True)
        self._car_combo.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self._car_combo.setStyleSheet(self._field_qss("combo"))
        self._car_combo.lineEdit().setPlaceholderText("Nom du transporteur…")

        bl.addLayout(self._frm_row("Transporteur :", self._car_combo))
        lo.addWidget(box)

        b = QPushButton("Rapport transporteurs (PDF)")
        b.setObjectName("primaryBtn")
        b.clicked.connect(self._gen_carrier)
        lo.addWidget(b)

        self._car_preview, self._car_dl_btn, _car_prev_hdr = self._make_preview_block()
        lo.addWidget(_car_prev_hdr)
        lo.addWidget(self._car_preview)
        lo.addWidget(self._car_dl_btn)
        lo.addStretch()
        return w

    def _make_compliance_panel(self):
        w, lo = self._panel_shell("Conformité RSE")
        box, bl = self._box(C["border"])

        self._co_start = QDateEdit()
        self._co_end = QDateEdit()
        for d in (self._co_start, self._co_end):
            d.setCalendarPopup(True)
            d.setDate(datetime.today().date())
            d.setMaximumWidth(150)
            d.setStyleSheet(self._field_qss("date"))

        bl.addLayout(self._frm_row("Début :", self._co_start))
        bl.addLayout(self._frm_row("Fin :", self._co_end))
        lo.addWidget(box)

        b = QPushButton("Rapport RSE (PDF)")
        b.setObjectName("primaryBtn")
        b.clicked.connect(self._gen_rse)
        lo.addWidget(b)

        self._co_preview, self._co_dl_btn, _co_prev_hdr = self._make_preview_block()
        lo.addWidget(_co_prev_hdr)
        lo.addWidget(self._co_preview)
        lo.addWidget(self._co_dl_btn)
        lo.addStretch()
        return w

    def _make_legal_panel(self):
        w, lo = self._panel_shell("Documents légaux & transport")

        _grp_qss = (
            f"QGroupBox{{border:1px solid {C['border']};border-radius:8px;"
            "margin-top:18px;padding:10px 10px 10px 10px;}}"
            f"QGroupBox::title{{subcontrol-origin:margin;subcontrol-position:top left;"
            f"top:-2px;left:10px;padding:0 4px;color:{C['accent']};"
            f"background:{C['panel']};}}"
        )
        _combo_qss = (
            f"QComboBox{{background:{C['bg']};color:{C['text']};"
            f"border:1px solid {C['border']};border-radius:4px;padding:4px;}}"
            f"QComboBox::drop-down{{border:none;}}"
            f"QComboBox QAbstractItemView{{background:{C['bg']};color:{C['text']};}}"
        )

        # ── Groupe 1 : éditeur CGU / Confidentialité ─────────────────────
        g = QGroupBox("Modèles généraux")
        g.setStyleSheet(_grp_qss)
        g_lo = QVBoxLayout(g)
        g_lo.setContentsMargins(10, 14, 10, 10)
        g_lo.setSpacing(10)

        self._leg_type = QComboBox()
        self._leg_type.addItem("CGU (synthèse)", "terms")
        self._leg_type.addItem("Confidentialité (synthèse)", "privacy")
        self._leg_type.setStyleSheet(_combo_qss)
        self._leg_type.setMaximumWidth(240)
        g_lo.addLayout(self._frm_row("Type :", self._leg_type, 80))

        note_edit = QLabel("Éditez le contenu ci-dessous, puis faites Aperçu ou Télécharger PDF.")
        note_edit.setStyleSheet(f"color:{C['muted']};font-size:10px;")
        note_edit.setWordWrap(True)
        g_lo.addWidget(note_edit)

        self._legal_editor = QTextEdit()
        self._legal_editor.setMinimumHeight(260)
        self._legal_editor.setStyleSheet(
            f"QTextEdit{{background:{C['panel']};color:{C['text']};"
            f"border:1px solid {C['border']};border-radius:4px;padding:8px;"
            f"font-family:Consolas,monospace;font-size:11px;}}"
        )
        g_lo.addWidget(self._legal_editor)

        btn_row = QHBoxLayout()
        btn_row.setSpacing(8)
        b_prev = QPushButton("Aperçu")
        b_prev.setObjectName("secondaryBtn")
        b_prev.clicked.connect(self._legal_preview)
        b_dl = QPushButton("Télécharger PDF")
        b_dl.setObjectName("primaryBtn")
        b_dl.clicked.connect(self._legal_download)
        btn_row.addWidget(b_prev)
        btn_row.addWidget(b_dl)
        btn_row.addStretch()
        g_lo.addLayout(btn_row)
        lo.addWidget(g)

        self._leg_type.currentIndexChanged.connect(self._on_legal_type_changed)
        self._on_legal_type_changed()  # pré-remplir dès l'ouverture

        # ── Groupe 2 : documents par commande / tournée ───────────────────
        g2 = QGroupBox("Documents commande / tournée")
        g2.setStyleSheet(_grp_qss)
        g2_lo = QVBoxLayout(g2)
        g2_lo.setContentsMargins(10, 14, 10, 10)
        g2_lo.setSpacing(10)

        note2 = QLabel("Cherchez par référence, nom client ou date — sélectionnez dans la liste.")
        note2.setStyleSheet(f"color:{C['muted']};font-size:10px;")
        note2.setWordWrap(True)
        g2_lo.addWidget(note2)

        self._leg_order_combo = QComboBox()
        self._leg_order_combo.setEditable(True)
        self._leg_order_combo.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self._leg_order_combo.setStyleSheet(_combo_qss)
        self._leg_order_combo.lineEdit().setPlaceholderText("Référence commande ou nom client…")

        self._leg_route_combo = QComboBox()
        self._leg_route_combo.setEditable(True)
        self._leg_route_combo.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self._leg_route_combo.setStyleSheet(_combo_qss)
        self._leg_route_combo.lineEdit().setPlaceholderText("Date ou immatriculation véhicule…")

        g2_lo.addLayout(self._frm_row("Commande (BL/CMR/ADR) :", self._leg_order_combo, 200))
        g2_lo.addLayout(self._frm_row("Tournée / Manifeste :", self._leg_route_combo, 200))
        lo.addWidget(g2)

        row = QHBoxLayout()
        row.setSpacing(6)
        for txt, slot in [
            ("BL", self._gen_legal_bl),
            ("CMR", self._gen_legal_cmr),
            ("ADR", self._gen_legal_adr),
            ("Manifeste", self._gen_legal_manifest),
        ]:
            bx = QPushButton(txt)
            bx.setObjectName("secondaryBtn")
            bx.clicked.connect(slot)
            row.addWidget(bx)
        lo.addLayout(row)

        self._leg_preview, self._leg_dl_btn, _leg_prev_hdr = self._make_preview_block()
        lo.addWidget(_leg_prev_hdr)
        lo.addWidget(self._leg_preview)
        lo.addWidget(self._leg_dl_btn)
        lo.addStretch()
        return w

    # ── Éditeur CGU / Confidentialité ─────────────────────────────────────────

    def _on_legal_type_changed(self):
        self._legal_editor.setPlainText(self._legal_get_text(self._leg_type.currentData()))

    def _legal_get_text(self, dtype: str) -> str:
        co: dict = {}
        try:
            sp = os.path.normpath(os.path.join(os.path.dirname(__file__), "..", "..", "settings.json"))
            with open(sp, encoding="utf-8") as f:
                co = json.load(f).get("company", {})
        except Exception:
            pass
        name  = co.get("name")    or "CityPulse Logistics"
        addr  = co.get("address") or "—"
        phone = co.get("phone")   or "—"
        email = co.get("email")   or "—"
        today = datetime.now().strftime("%d/%m/%Y")

        if dtype == "terms":
            return (
                f"CONDITIONS GÉNÉRALES D'UTILISATION\n"
                f"{name} — Version du {today}\n"
                f"Siège : {addr} | Tél : {phone} | Email : {email}\n\n"
                f"Article 1 — Objet\n"
                f"Le présent document définit les conditions d'utilisation du logiciel de gestion "
                f"et d'optimisation de tournées exploité par {name}. Il s'applique à tout utilisateur "
                f"ayant accès au système (opérateur, planificateur, administrateur).\n\n"
                f"Article 2 — Accès au système\n"
                f"L'accès est réservé aux personnes habilitées. Chaque utilisateur dispose d'identifiants "
                f"personnels (login + mot de passe hashé). Le partage de compte est interdit. "
                f"Toute compromission doit être signalée immédiatement à l'administrateur.\n\n"
                f"Article 3 — Utilisation acceptable\n"
                f"Le logiciel est utilisé exclusivement dans le cadre professionnel de l'entreprise. "
                f"Toute utilisation à des fins personnelles, concurrentielles ou frauduleuses est interdite.\n\n"
                f"Article 4 — Données saisies et responsabilité\n"
                f"Les données saisies (clients, commandes, coordonnées GPS, véhicules, tournées) restent "
                f"la propriété de {name}. L'utilisateur est responsable de l'exactitude des informations.\n\n"
                f"Article 5 — Conservation des données\n"
                f"Les données opérationnelles sont conservées en base SQLite locale. "
                f"Des sauvegardes sont disponibles depuis Paramètres → Sauvegarde.\n\n"
                f"Article 6 — Droits des utilisateurs\n"
                f"Tout utilisateur peut demander consultation, rectification ou suppression de ses données "
                f"personnelles auprès de l'administrateur système.\n\n"
                f"Article 7 — Modifications\n"
                f"Ces conditions peuvent être mises à jour par l'administrateur de {name}. "
                f"Les utilisateurs seront informés de toute modification substantielle.\n\n"
                f"Article 8 — Contact\n"
                f"{name} — {addr}\nTél : {phone} | Email : {email}"
            )
        else:
            return (
                f"POLITIQUE DE PROTECTION DES DONNÉES PERSONNELLES\n"
                f"{name} — Version du {today}\n"
                f"Siège : {addr} | Tél : {phone} | Email : {email}\n\n"
                f"Article 1 — Responsable du traitement\n"
                f"Le responsable du traitement est {name}, dont le siège est situé à {addr}. "
                f"Contact : {email} — Tél : {phone}.\n\n"
                f"Article 2 — Données collectées\n"
                f"• Clients : nom, adresse de livraison, coordonnées GPS, téléphone, email\n"
                f"• Chauffeurs : nom, prénom, numéro de permis, qualifications, photo\n"
                f"• Opérationnel : commandes, tournées, horaires, distances, coûts, CO₂\n"
                f"• Système : logs d'audit, sessions utilisateurs, horodatages\n\n"
                f"Article 3 — Finalités du traitement\n"
                f"• Planification et optimisation des tournées de livraison\n"
                f"• Gestion de la flotte de véhicules et des ressources humaines\n"
                f"• Suivi opérationnel, reporting et facturation\n"
                f"• Conformité réglementaire (CE 561/2006, ADR, ZFE)\n\n"
                f"Article 4 — Base légale\n"
                f"Le traitement est fondé sur l'intérêt légitime de l'entreprise pour la gestion "
                f"de son activité de transport et les obligations légales applicables.\n\n"
                f"Article 5 — Durée de conservation\n"
                f"• Données clients et commandes : relation commerciale + 5 ans\n"
                f"• Données chauffeurs : durée du contrat + 5 ans\n"
                f"• Logs d'audit : 3 ans\n"
                f"• Données de géolocalisation (tournées) : 2 ans\n\n"
                f"Article 6 — Droits des personnes concernées\n"
                f"Accès, rectification, effacement, opposition, portabilité. "
                f"Pour exercer ces droits : {email}.\n\n"
                f"Article 7 — Sécurité des données\n"
                f"Données stockées localement (SQLite). Mots de passe hashés SHA-256. "
                f"Clés API dans le trousseau système (OS keyring). Sauvegardes recommandées.\n\n"
                f"Article 8 — Contact DPO\n"
                f"{name} — {addr}\nEmail : {email} | Tél : {phone}"
            )

    def _legal_preview(self):
        if not HAS_WEB:
            return
        text = self._legal_editor.toPlainText()
        lines = text.split("\n")
        html_body = ""
        for i, line in enumerate(lines):
            s = line.strip()
            if not s:
                html_body += "<br/>"
            elif i == 0:
                html_body += f"<h2 style='color:#1E3A5F;margin:0 0 4px'>{s}</h2>"
            elif i == 1:
                html_body += f"<p style='color:#3B82F6;font-size:12px;margin:0 0 2px'>{s}</p>"
            elif i == 2:
                html_body += (f"<p style='color:#6B7280;font-size:11px;margin:0 0 14px'>{s}</p>"
                              f"<hr style='border:none;border-top:1px solid #e5e7eb;margin-bottom:12px'/>")
            elif s.startswith("Article "):
                html_body += f"<h4 style='color:#1E3A5F;margin:14px 0 4px'>{s}</h4>"
            elif s.startswith("•"):
                html_body += f"<p style='margin:2px 0 2px 16px;font-size:12px'>{s}</p>"
            else:
                html_body += f"<p style='margin:3px 0;font-size:12px'>{s}</p>"
        html = (
            "<html><body style='font-family:Arial,sans-serif;padding:28px;"
            f"color:#111827;max-width:800px'>{html_body}</body></html>"
        )
        self._leg_preview.setHtml(html)

    def _legal_download(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        text = self._legal_editor.toPlainText().strip()
        if not text:
            QMessageBox.warning(self, "Contenu vide", "L'éditeur est vide.")
            return
        dtype = self._leg_type.currentData()
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"legal_{dtype}.pdf"
        self._run_worker(lambda: self._service.generate_legal_from_text(text, path),
                         preview=self._leg_preview, dl_btn=self._leg_dl_btn)

    # ── Combos searchables commandes / tournées ───────────────────────────────

    def _load_order_options(self):
        try:
            conn = get_connection()
            rows = conn.execute(
                """SELECT o.id, o.reference, c.name AS client_name
                   FROM orders o
                   LEFT JOIN clients c ON c.id = o.client_id
                   WHERE o.archived = 0
                   ORDER BY o.id DESC LIMIT 500"""
            ).fetchall()
            conn.close()
            self._leg_order_combo.blockSignals(True)
            self._leg_order_combo.clear()
            for row in rows:
                label = f"{row['reference'] or 'CMD'} — {row['client_name'] or '?'} (ID={row['id']})"
                self._leg_order_combo.addItem(label, row["id"])
            items = [self._leg_order_combo.itemText(i) for i in range(self._leg_order_combo.count())]
            cmp = QCompleter(items, self._leg_order_combo)
            cmp.setFilterMode(Qt.MatchFlag.MatchContains)
            cmp.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
            self._leg_order_combo.setCompleter(cmp)
            self._leg_order_combo.blockSignals(False)
        except Exception:
            pass

    def _load_route_options(self):
        try:
            conn = get_connection()
            rows = conn.execute(
                """SELECT r.id, r.planned_date, v.registration
                   FROM routes r
                   LEFT JOIN vehicles v ON v.id = r.vehicle_id
                   ORDER BY r.id DESC LIMIT 300"""
            ).fetchall()
            conn.close()
            self._leg_route_combo.blockSignals(True)
            self._leg_route_combo.clear()
            for row in rows:
                label = (f"{row['planned_date'] or '?'} — "
                         f"{row['registration'] or 'Véhicule'} (ID={row['id']})")
                self._leg_route_combo.addItem(label, row["id"])
            items = [self._leg_route_combo.itemText(i) for i in range(self._leg_route_combo.count())]
            cmp = QCompleter(items, self._leg_route_combo)
            cmp.setFilterMode(Qt.MatchFlag.MatchContains)
            cmp.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
            self._leg_route_combo.setCompleter(cmp)
            self._leg_route_combo.blockSignals(False)
        except Exception:
            pass

    def _get_order_id(self) -> int | None:
        data = self._leg_order_combo.currentData()
        if data:
            return int(data)
        m = re.search(r"ID=(\d+)", self._leg_order_combo.currentText())
        return int(m.group(1)) if m else None

    def _get_route_id(self) -> int | None:
        data = self._leg_route_combo.currentData()
        if data:
            return int(data)
        m = re.search(r"ID=(\d+)", self._leg_route_combo.currentText())
        return int(m.group(1)) if m else None

    # ── Loaders et helpers pour les autres panels ─────────────────────────────

    def _load_op_route_options(self):
        try:
            conn = get_connection()
            rows = conn.execute(
                """SELECT r.id, r.planned_date, v.registration
                   FROM routes r LEFT JOIN vehicles v ON v.id = r.vehicle_id
                   ORDER BY r.id DESC LIMIT 300"""
            ).fetchall()
            conn.close()
            self._op_route_combo.blockSignals(True)
            self._op_route_combo.clear()
            for row in rows:
                label = (f"{row['planned_date'] or '?'} — "
                         f"{row['registration'] or 'Véhicule'} (ID={row['id']})")
                self._op_route_combo.addItem(label, row["id"])
            cmp = QCompleter(
                [self._op_route_combo.itemText(i) for i in range(self._op_route_combo.count())],
                self._op_route_combo)
            cmp.setFilterMode(Qt.MatchFlag.MatchContains)
            cmp.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
            self._op_route_combo.setCompleter(cmp)
            self._op_route_combo.blockSignals(False)
        except Exception:
            pass

    def _load_algo_options(self):
        try:
            conn = get_connection()
            rows = conn.execute(
                """SELECT id, algorithm, total_distance, created_at
                   FROM algo_results ORDER BY id DESC LIMIT 200"""
            ).fetchall()
            conn.close()
            self._an_algo_pick.blockSignals(True)
            self._an_algo_pick.clear()
            self._an_algo_pick.addItem("— Sélectionner —", None)
            for row in rows:
                dist = f"{float(row['total_distance'] or 0):.0f} km" if row['total_distance'] else "?"
                dt = (row['created_at'] or "")[:10]
                label = f"{row['algorithm'] or '?'} — {dist} — {dt} (ID={row['id']})"
                self._an_algo_pick.addItem(label, row["id"])
            cmp = QCompleter(
                [self._an_algo_pick.itemText(i) for i in range(self._an_algo_pick.count())],
                self._an_algo_pick)
            cmp.setFilterMode(Qt.MatchFlag.MatchContains)
            cmp.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
            self._an_algo_pick.setCompleter(cmp)
            self._an_algo_pick.blockSignals(False)
        except Exception:
            pass

    def _on_algo_pick(self, idx: int):
        rid = self._an_algo_pick.itemData(idx)
        if not rid:
            return
        # Vérifier doublon dans la liste
        for i in range(self._an_algo_list.count()):
            if self._an_algo_list.item(i).data(Qt.ItemDataRole.UserRole) == rid:
                return
        label = self._an_algo_pick.itemText(idx)
        it = QListWidgetItem(label)
        it.setData(Qt.ItemDataRole.UserRole, rid)
        self._an_algo_list.addItem(it)

    def _algo_list_remove(self):
        for it in self._an_algo_list.selectedItems():
            self._an_algo_list.takeItem(self._an_algo_list.row(it))

    def _load_client_options(self):
        try:
            conn = get_connection()
            rows = conn.execute(
                """SELECT id, name, company_name FROM clients
                   WHERE archived=0 ORDER BY name ASC LIMIT 500"""
            ).fetchall()
            conn.close()
            self._cl_combo.blockSignals(True)
            self._cl_combo.clear()
            for row in rows:
                co = row["company_name"] or ""
                label = f"{row['name']}" + (f" ({co})" if co else "") + f" — ID={row['id']}"
                self._cl_combo.addItem(label, row["id"])
            cmp = QCompleter(
                [self._cl_combo.itemText(i) for i in range(self._cl_combo.count())],
                self._cl_combo)
            cmp.setFilterMode(Qt.MatchFlag.MatchContains)
            cmp.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
            self._cl_combo.setCompleter(cmp)
            self._cl_combo.blockSignals(False)
        except Exception:
            pass

    def _load_carrier_options(self):
        try:
            conn = get_connection()
            rows = conn.execute(
                "SELECT id, name FROM carriers ORDER BY name ASC LIMIT 200"
            ).fetchall()
            conn.close()
            self._car_combo.blockSignals(True)
            self._car_combo.clear()
            self._car_combo.addItem("Tous les transporteurs", None)
            for row in rows:
                self._car_combo.addItem(f"{row['name']} (ID={row['id']})", row["id"])
            cmp = QCompleter(
                [self._car_combo.itemText(i) for i in range(self._car_combo.count())],
                self._car_combo)
            cmp.setFilterMode(Qt.MatchFlag.MatchContains)
            cmp.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
            self._car_combo.setCompleter(cmp)
            self._car_combo.blockSignals(False)
        except Exception:
            pass

    def _get_op_route_id(self) -> int | None:
        data = self._op_route_combo.currentData()
        if data:
            return int(data)
        m = re.search(r"ID=(\d+)", self._op_route_combo.currentText())
        return int(m.group(1)) if m else None

    def _get_client_id(self) -> int | None:
        data = self._cl_combo.currentData()
        if data:
            return int(data)
        m = re.search(r"ID=(\d+)", self._cl_combo.currentText())
        return int(m.group(1)) if m else None

    def _get_carrier_id(self) -> int | None:
        data = self._car_combo.currentData()
        return int(data) if data else None

    def _make_exports_panel(self):
        w, lo = self._panel_shell("Exports massifs")
        r = QHBoxLayout()
        b1 = QPushButton("Excel multi-feuilles")
        b1.setObjectName("primaryBtn")
        b1.clicked.connect(self._gen_excel)
        b2 = QPushButton("Snapshot JSON complet")
        b2.setObjectName("secondaryBtn")
        b2.clicked.connect(self._gen_snapshot)
        r.addWidget(b1)
        r.addWidget(b2)
        lo.addLayout(r)
        note = QLabel("Inclut Clients, Véhicules, Chauffeurs, Commandes, Tournées, Journal (5000 lignes max).")
        note.setWordWrap(True)
        note.setStyleSheet(f"color:{C['muted']};font-size:10px;")
        lo.addWidget(note)

        self._exp_preview, self._exp_dl_btn, _exp_prev_hdr = self._make_preview_block()
        lo.addWidget(_exp_prev_hdr)
        lo.addWidget(self._exp_preview)
        lo.addWidget(self._exp_dl_btn)
        lo.addStretch()
        return w

    def _on_category_changed(self, cur: QListWidgetItem | None, _prev):
        if not cur:
            return
        key = cur.data(Qt.ItemDataRole.UserRole)
        idx = {
            "ops": 0,
            "analytics": 1,
            "clients": 2,
            "carriers": 3,
            "compliance": 4,
            "legal": 5,
            "exports": 6,
        }.get(key, 0)
        self._stack.setCurrentIndex(idx)

    def _run_worker(self, fn, *args, preview=None, dl_btn=None, **kwargs):
        if self._worker and self._worker.isRunning():
            show_toast(self.window(), "Une génération est déjà en cours.", "info")
            return
        self._active_preview = preview
        self._active_dl_btn = dl_btn
        self._overlay.setGeometry(self.rect())
        self._overlay.start("Génération du rapport…")
        self._worker = _ReportWorker(fn, *args, **kwargs)
        self._worker.finished_ok.connect(self._on_report_done)
        self._worker.failed.connect(self._on_report_err)
        self._worker.start()

    def _on_report_done(self, path: str):
        self._overlay.stop()
        self._last_pdf_path = path
        preview = getattr(self, "_active_preview", None)
        dl_btn = getattr(self, "_active_dl_btn", None)
        if preview:
            self._load_preview(preview, path)
        low = path.lower()
        is_previewable = low.endswith(".pdf") or low.endswith(".xlsx") or low.endswith(".json")
        if dl_btn:
            if low.endswith(".xlsx"):
                dl_btn.setText("⬇  Télécharger l'Excel")
            elif low.endswith(".json"):
                dl_btn.setText("⬇  Télécharger le JSON")
            else:
                dl_btn.setText("⬇  Télécharger le rapport")
            dl_btn.setVisible(is_previewable)
        show_toast(self.window(), "Fichier prêt — cliquez Télécharger pour enregistrer.", "success")
        self.refresh_data()

    def _download_last_report(self):
        if not self._last_pdf_path or not os.path.isfile(self._last_pdf_path):
            show_toast(self.window(), "Aucun rapport disponible.", "error")
            return
        low = self._last_pdf_path.lower()
        if low.endswith(".xlsx"):
            filt = "Excel (*.xlsx)"
        elif low.endswith(".json"):
            filt = "JSON (*.json)"
        else:
            filt = "PDF (*.pdf)"
        dest, _ = QFileDialog.getSaveFileName(
            self, "Enregistrer le fichier", self._last_pdf_default_name, filt
        )
        if not dest:
            return
        try:
            shutil.copy(self._last_pdf_path, dest)
            show_toast(self.window(), f"Enregistré : {dest}", "success")
        except Exception as exc:
            QMessageBox.warning(self, "Erreur", str(exc))

    def _on_report_err(self, msg: str):
        self._overlay.stop()
        QMessageBox.warning(self, "Rapport", msg)
        log_action("REPORT_ERROR", msg)

    def _load_preview(self, preview: "_PreviewPane", path: str):
        if not path or not os.path.isfile(path):
            return
        low = path.lower()
        if low.endswith(".pdf"):
            preview.show_pdf(path)
            return
        if not HAS_WEB:
            return
        if low.endswith(".json"):
            try:
                with open(path, encoding="utf-8") as f:
                    data = json.load(f)
                lines = []
                if isinstance(data, dict):
                    for table, records in list(data.items())[:10]:
                        count = len(records) if isinstance(records, list) else "?"
                        lines.append(
                            f"<tr><td style='color:{C['accent']};padding:4px 10px;font-weight:600'>"
                            f"{table}</td>"
                            f"<td style='color:{C['text']};padding:4px 10px'>{count} enregistrement(s)</td></tr>"
                        )
                    rows_html = "".join(lines)
                    total = sum(len(v) for v in data.values() if isinstance(v, list))
                    html = (
                        f"<html><body style='background:{C['bg']};color:{C['text']};"
                        f"font-family:monospace;font-size:12px;margin:0;padding:12px'>"
                        f"<p style='color:{C['muted']};font-size:11px;margin:0 0 10px'>"
                        f"Snapshot JSON — {os.path.basename(path)} — {total} enregistrements au total</p>"
                        f"<table style='border-collapse:collapse;width:100%'>{rows_html}</table>"
                        f"</body></html>"
                    )
                else:
                    snippet = json.dumps(data, ensure_ascii=False, indent=2)[:3000]
                    html = (
                        f"<html><body style='background:{C['bg']};color:{C['text']};"
                        f"font-family:monospace;font-size:11px;margin:0;padding:12px'>"
                        f"<pre>{snippet}</pre></body></html>"
                    )
                preview.setHtml(html)
            except Exception:
                preview.setHtml(
                    f"<body style='background:{C['bg']};color:{C['muted']};"
                    f"font-family:sans-serif;padding:24px'>"
                    f"<p>Impossible de lire le fichier JSON.</p></body>"
                )
            return
        if low.endswith(".xlsx"):
            if OPENPYXL_OK:
                try:
                    import openpyxl as _xl
                    wb = _xl.load_workbook(path, read_only=True, data_only=True)
                    ws = wb.active
                    rows = list(ws.iter_rows(max_row=60, values_only=True))
                    wb.close()
                    html_rows = ""
                    for i, row in enumerate(rows):
                        bg = "#112240" if i % 2 == 0 else "#0D1B2A"
                        tds = "".join(
                            f"<td style='padding:4px 8px;border:1px solid #1E3A5F;"
                            f"white-space:nowrap;max-width:220px;overflow:hidden;text-overflow:ellipsis'>"
                            f"{str(c) if c is not None else ''}</td>"
                            for c in row
                        )
                        html_rows += f"<tr style='background:{bg}'>{tds}</tr>"
                    fname = os.path.basename(path)
                    html = (
                        "<html><body style='background:#0D1B2A;color:#E8F4FD;"
                        "font-family:monospace;font-size:12px;margin:0;padding:12px'>"
                        f"<p style='color:#8899AA;font-size:11px;margin:0 0 8px'>"
                        f"Aperçu — {fname} (60 premières lignes)</p>"
                        "<div style='overflow-x:auto'>"
                        "<table style='border-collapse:collapse;width:100%'>"
                        f"{html_rows}</table></div></body></html>"
                    )
                    preview.setHtml(html)
                except Exception:
                    preview.setHtml(
                        f"<body style='background:#0D1B2A;color:#8899AA;"
                        f"font-family:sans-serif;padding:24px'>"
                        f"<p>Impossible de lire le fichier.</p>"
                        f"<p>Fichier : <b>{path}</b></p></body>"
                    )
            else:
                preview.setHtml(
                    f"<body style='background:#0D1B2A;color:#8899AA;"
                    f"font-family:sans-serif;padding:24px'>"
                    f"<p>Installez <b>openpyxl</b> pour l'aperçu Excel.</p>"
                    f"<p>Fichier : <b>{path}</b></p></body>"
                )

    def _save_dialog(self, title: str, default: str, filt: str) -> str | None:
        path, _ = QFileDialog.getSaveFileName(self, title, default, filt)
        return path or None

    def _gen_roadbook(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        rid = self._get_op_route_id()
        if not rid:
            QMessageBox.warning(self, "Tournée", "Sélectionnez une tournée dans la liste.")
            return
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"roadbook_{rid}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_driver_roadbook(rid, path, lang=lang),
                         preview=self._ops_preview, dl_btn=self._ops_dl_btn)

    def _gen_fleet_daily(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        d = self._op_date.date().toString("yyyy-MM-dd")
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"flotte_{d}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_fleet_daily_report(d, path, lang=lang),
                         preview=self._ops_preview, dl_btn=self._ops_dl_btn)

    def _gen_kpi(self):
        s = self._an_kpi_start.date().toString("yyyy-MM-dd")
        e = self._an_kpi_end.date().toString("yyyy-MM-dd")
        fmt = self._an_kpi_fmt.currentText()
        ext = "xlsx" if fmt == "xlsx" else "pdf"
        if fmt == "xlsx" and not OPENPYXL_OK:
            QMessageBox.warning(self, "Dépendance", "Installez openpyxl.")
            return
        if ext == "pdf" and not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        if fmt == "xlsx":
            path = self._save_dialog("Rapport KPI", f"kpi_{s}_{e}.xlsx", "Excel (*.xlsx)")
            if not path:
                return
        else:
            fd, path = tempfile.mkstemp(suffix=".pdf")
            os.close(fd)
            self._last_pdf_default_name = f"kpi_{s}_{e}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_kpi_report(s, e, path, fmt=fmt, lang=lang),
                         preview=self._an_preview, dl_btn=self._an_dl_btn)

    def _gen_algo_cmp(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        ids = [
            self._an_algo_list.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self._an_algo_list.count())
        ]
        if not ids:
            QMessageBox.warning(self, "Paramètres",
                                "Ajoutez au moins un résultat à comparer via la liste déroulante.")
            return
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = "comparaison_algos.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_algo_comparison_report(ids, path, lang=lang),
                         preview=self._an_preview, dl_btn=self._an_dl_btn)

    def _gen_drv_perf(self):
        days = self._an_drv_days.value()
        fmt = self._an_drv_fmt.currentText()
        ext = "xlsx" if fmt == "xlsx" else "pdf"
        if fmt == "xlsx" and not OPENPYXL_OK:
            QMessageBox.warning(self, "Dépendance", "Installez openpyxl.")
            return
        if ext == "pdf" and not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        if fmt == "xlsx":
            path = self._save_dialog("Performance chauffeurs", f"chauffeurs_{days}j.xlsx", "Excel (*.xlsx)")
            if not path:
                return
        else:
            fd, path = tempfile.mkstemp(suffix=".pdf")
            os.close(fd)
            self._last_pdf_default_name = f"chauffeurs_{days}j.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_driver_performance_report(days, path, fmt=fmt, lang=lang),
                         preview=self._an_preview, dl_btn=self._an_dl_btn)

    def _gen_client(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        cid = self._get_client_id()
        if not cid:
            QMessageBox.warning(self, "Client", "Sélectionnez un client dans la liste.")
            return
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"client_{cid}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_client_report(cid, path, lang=lang),
                         preview=self._cl_preview, dl_btn=self._cl_dl_btn)

    def _gen_carrier(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        cid = self._get_carrier_id()
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = "transporteurs.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(
            lambda: self._service.generate_carrier_report(carrier_id=cid, output_path=path, lang=lang),
            preview=self._car_preview, dl_btn=self._car_dl_btn
        )

    def _gen_rse(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        s = self._co_start.date().toString("yyyy-MM-dd")
        e = self._co_end.date().toString("yyyy-MM-dd")
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"rse_{s}_{e}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_rse_compliance_report(s, e, path, lang=lang),
                         preview=self._co_preview, dl_btn=self._co_dl_btn)

    def _gen_legal(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        dtype = self._leg_type.currentData()
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"legal_{dtype}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_legal_notice_pdf(path, doc_type=dtype, lang=lang),
                         preview=self._leg_preview, dl_btn=self._leg_dl_btn)

    def _gen_legal_bl(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        oid = self._get_order_id()
        if not oid:
            QMessageBox.warning(self, "Commande", "Sélectionnez une commande dans la liste.")
            return
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"BL_{oid}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_delivery_note(oid, path, lang=lang),
                         preview=self._leg_preview, dl_btn=self._leg_dl_btn)

    def _gen_legal_cmr(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        oid = self._get_order_id()
        if not oid:
            QMessageBox.warning(self, "Commande", "Sélectionnez une commande dans la liste.")
            return
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"CMR_{oid}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_cmr(oid, path, lang=lang),
                         preview=self._leg_preview, dl_btn=self._leg_dl_btn)

    def _gen_legal_adr(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        oid = self._get_order_id()
        if not oid:
            QMessageBox.warning(self, "Commande", "Sélectionnez une commande dans la liste.")
            return
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"ADR_{oid}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_adr_document(oid, path, lang=lang),
                         preview=self._leg_preview, dl_btn=self._leg_dl_btn)

    def _gen_legal_manifest(self):
        if not REPORTLAB_OK:
            QMessageBox.warning(self, "Dépendance", "Installez reportlab.")
            return
        rid = self._get_route_id()
        if not rid:
            QMessageBox.warning(self, "Tournée", "Sélectionnez une tournée dans la liste.")
            return
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        self._last_pdf_default_name = f"manifeste_route_{rid}.pdf"
        lang = self._report_lang_combo.currentData()
        self._run_worker(lambda: self._service.generate_load_manifest(rid, path, lang=lang),
                         preview=self._leg_preview, dl_btn=self._leg_dl_btn)

    def _gen_excel(self):
        if not OPENPYXL_OK:
            QMessageBox.warning(self, "Dépendance", "Installez openpyxl.")
            return
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        self._last_pdf_default_name = "citypulse_export.xlsx"
        self._run_worker(lambda: self._service.export_to_excel(path),
                         preview=self._exp_preview, dl_btn=self._exp_dl_btn)

    def _gen_snapshot(self):
        fd, path = tempfile.mkstemp(suffix=".json")
        os.close(fd)
        self._last_pdf_default_name = "citypulse_snapshot.json"
        self._run_worker(lambda: self._service.generate_full_snapshot(path),
                         preview=self._exp_preview, dl_btn=self._exp_dl_btn)

    def _settings_path(self) -> str:
        return os.path.normpath(
            os.path.join(os.path.dirname(__file__), "..", "..", "settings.json")
        )

    def _load_sched_settings(self):
        try:
            with open(self._settings_path(), encoding="utf-8") as f:
                cfg = json.load(f)
            s = cfg.get("reports", {}).get("schedule", {})
            self._sched_enable.blockSignals(True)
            self._sched_dir.blockSignals(True)
            self._sched_time.blockSignals(True)
            self._sched_enable.setChecked(bool(s.get("enabled", False)))
            self._sched_dir.setText(s.get("directory", ""))
            h, m = s.get("hour", 8), s.get("minute", 0)
            self._sched_time.setTime(QTime(int(h), int(m)))
            self._sched_enable.blockSignals(False)
            self._sched_dir.blockSignals(False)
            self._sched_time.blockSignals(False)
        except Exception:
            pass

    def _save_sched_settings(self):
        try:
            sp = self._settings_path()
            try:
                with open(sp, encoding="utf-8") as f:
                    cfg = json.load(f)
            except Exception:
                cfg = {}
            cfg.setdefault("reports", {})["schedule"] = {
                "enabled": self._sched_enable.isChecked(),
                "directory": self._sched_dir.text().strip(),
                "hour": self._sched_time.time().hour(),
                "minute": self._sched_time.time().minute(),
            }
            with open(sp, "w", encoding="utf-8") as f:
                json.dump(cfg, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def _pick_sched_dir(self):
        d = QFileDialog.getExistingDirectory(self, "Dossier exports planifiés")
        if d:
            self._sched_dir.setText(d)

    def _on_schedule_tick(self):
        if not self._sched_enable.isChecked():
            return
        t = self._sched_time.time()
        now = datetime.now()
        if now.hour != t.hour() or now.minute != t.minute():
            return
        today = now.date().isoformat()
        if self._sched_last_kpi_day == today:
            return
        base = self._sched_dir.text().strip() or os.path.join(os.path.expanduser("~"), "Documents", "CityPulseReports")
        try:
            os.makedirs(base, exist_ok=True)
        except OSError:
            return
        if not REPORTLAB_OK:
            return
        end = date.today()
        start = end - timedelta(days=7)
        path = os.path.join(base, f"kpi_auto_{end.isoformat()}.pdf")
        try:
            self._service.generate_kpi_report(start.isoformat(), end.isoformat(), path, fmt="pdf")
            self._sched_last_kpi_day = today
            log_action("REPORT_SCHEDULED_KPI", path)
        except Exception:
            pass

    def _open_history_file(self, item: QListWidgetItem):
        p = item.data(Qt.ItemDataRole.UserRole)
        if p and os.path.isfile(p):
            preview = getattr(self, "_active_preview", None)
            if preview:
                self._load_preview(preview, p)

    def retranslate_ui(self, lang: str):
        pass

    def refresh_data(self):
        self._load_sched_settings()
        self._load_order_options()
        self._load_route_options()
        self._load_op_route_options()
        self._load_algo_options()
        self._load_client_options()
        self._load_carrier_options()
        self._history_list.clear()
        try:
            conn = get_connection()
            rows = conn.execute(
                """SELECT report_type, file_path, generated_at, file_size_kb
                   FROM reports_history ORDER BY id DESC LIMIT 25"""
            ).fetchall()
            conn.close()
        except Exception:
            rows = []
        for r in rows:
            rd = {k: r[k] for k in r.keys()}
            txt = f"{str(rd.get('generated_at', ''))[:16]} | {rd.get('report_type')} | {rd.get('file_size_kb')} Ko"
            it = QListWidgetItem(txt)
            it.setData(Qt.ItemDataRole.UserRole, rd.get("file_path"))
            self._history_list.addItem(it)
